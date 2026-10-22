"""Pravi reference.docx sa stilovima po uputstvu fakulteta (forma-diplomski-rad-2025):
A4, margine 2,5 cm, Times New Roman 12, razmak 1, obostrano poravnanje, uvlacenje 12,7 mm,
naslovi boldovani 16/14 pt, poglavlja na novoj strani, brojevi stranica dole u sredini.
Pokretanje: python3 make_reference_docx.py reference.docx
"""
import subprocess, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

out = sys.argv[1] if len(sys.argv) > 1 else "reference.docx"
with open(out, "wb") as f:
    f.write(subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                           capture_output=True, check=True).stdout)
doc = Document(out)

def set_font(style, size, bold=None, italic=None):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None: style.font.bold = bold
    if italic is not None: style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), "Times New Roman")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None: del rfonts.attrib[qn(a)]

def para(style, align=None, indent=None, before=0, after=0, keep_next=None, page_break=None):
    pf = style.paragraph_format
    if align is not None: pf.alignment = align
    pf.first_line_indent = indent
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.line_spacing = 1.0
    if keep_next is not None: pf.keep_with_next = keep_next
    if page_break is not None: pf.page_break_before = page_break

st = doc.styles
def S(name):
    return next(x for x in st if x.name == name)
for name in ("Normal", "Body Text", "First Paragraph"):
    set_font(S(name), 12)
    para(S(name), WD_ALIGN_PARAGRAPH.JUSTIFY, Cm(1.27), after=0)
set_font(S("Compact"), 11)
para(S("Compact"), WD_ALIGN_PARAGRAPH.LEFT, Cm(0))
for name, size, pb in (("Heading 1", 16, True), ("Heading 2", 14, False), ("Heading 3", 12, False)):
    set_font(S(name), size, bold=True)
    para(S(name), WD_ALIGN_PARAGRAPH.LEFT, Cm(0), before=12, after=6, keep_next=True, page_break=pb)
set_font(S("Table Caption"), 12, bold=False, italic=False)
para(S("Table Caption"), WD_ALIGN_PARAGRAPH.LEFT, Cm(0), before=6, after=3, keep_next=True)
if any(x.name == "Bibliography" for x in st):
    set_font(S("Bibliography"), 12)
    para(S("Bibliography"), WD_ALIGN_PARAGRAPH.JUSTIFY, Cm(0), after=3)

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Cm(2.5))

# broj strane dole u sredini
p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    r = p.add_run(); set_font_run = r.font; set_font_run.name = "Times New Roman"; set_font_run.size = Pt(12)
    if t:
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), t); r._r.append(fc)
    else:
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)
doc.save(out)
print("OK", out)
